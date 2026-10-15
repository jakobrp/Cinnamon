import argparse
import csv
import statistics
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

try:
    import matplotlib.pyplot as plt
except ImportError:  # pragma: no cover - optional dependency
    plt = None

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - optional dependency
    load_workbook = None

DEFAULT_DATASET = "pima_diabetes"
BASE_DIR = Path(__file__).resolve().parent
DEFAULT_ROOT = BASE_DIR / "experiment_results" / DEFAULT_DATASET

SUMMARY_XLSX = "evaluation_summary.xlsx"
SUMMARY_CSV = "evaluation_summary.csv"

# Feste Reihenfolge fuer die Baseline; alles Weitere wird erkannt.
BASE_LABELS = ["real", "anon", "anon_synth"]

# Feste Farben der Baseline, DP-Varianten bekommen ihre Farbe ueber build_colors.
BASE_COLORS = {
    "real": "#333333",
    "anon": "#1f77b4",
    "anon_synth": "#2ca02c",
}

# Farbfamilien pro Verfahren: hell (kleines eps) -> dunkel (grosses eps).
METHOD_RAMPS = [
    ["#9ecae1", "#4292c6", "#08519c"],  # Blau
    ["#fcae91", "#fb6a4a", "#a50f15"],  # Rot
]


def split_label(label: str) -> Tuple[str, Optional[float]]:
    """'dpgan_eps0_1' -> ('dpgan', 0.1). Ohne eps: (label, None).
    Benötigt für Farbe und Pfeile
    """
    if "_eps" not in label:
        return label, None
    # Alles vor "_eps" ist das Verfahren, alles danach der eps-Wert.
    method, _, raw = label.partition("_eps")
    try:
        # Im Ordnernamen steht "_" statt ".", also 0_1 -> 0.1
        value = float(raw.replace("_", "."))
    except ValueError:
        value = None
    return method, value


def read_summary(summary_path: Path) -> List[Dict[str, Any]]:
    if summary_path.suffix.lower() == ".xlsx":
        return read_summary_xlsx(summary_path)
    return read_summary_csv(summary_path)


def read_summary_xlsx(summary_path: Path) -> List[Dict[str, Any]]:
    if load_workbook is None:
        raise RuntimeError("openpyxl is required to read Excel files. Install it with: pip install openpyxl")
    workbook = load_workbook(summary_path, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    data_rows: List[Dict[str, Any]] = []
    for row in rows[1:]:
        row_dict = {headers[i]: row[i] for i in range(len(headers))}
        data_rows.append(row_dict)
    return data_rows


def read_summary_csv(summary_path: Path) -> List[Dict[str, Any]]:
    data_rows: List[Dict[str, Any]] = []
    with summary_path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            data_rows.append(row)
    return data_rows


def discover_labels(rows: List[Dict[str, Any]]) -> List[str]:
    """Label aus den Spaltennamen ableiten: alles nach 'privacy_score_'.
    privacy_score_dpgan_eps1 --> dpgan_eps1
    """
    if not rows:
        return list(BASE_LABELS)
    prefix = "privacy_score_"
    found = [key[len(prefix):] for key in rows[0] if key.startswith(prefix)]
    # Erst die Baseline in fester Reihenfolge, danach die DP-Varianten alphabetisch.
    ordered = [label for label in BASE_LABELS if label in found]
    ordered += sorted(label for label in found if label not in BASE_LABELS)
    return ordered

# mithilfe von Claude generiert
def build_colors(labels: List[str]) -> Dict[str, str]:
    """Pro Verfahren eine Farbfamilie, innerhalb davon nach eps abgestuft.

    """
    colors = dict(BASE_COLORS)

    # Labels nach Verfahren gruppieren: {'dpgan': [(0.1, 'dpgan_eps0_1'), ...], 'pategan': [...]}
    by_method: Dict[str, List[Tuple[Optional[float], str]]] = {}
    for label in labels:
        # Nur für DP oder PATE GAN
        if label in BASE_COLORS:
            continue
        method, eps = split_label(label)
        by_method.setdefault(method, []).append((eps, label))

    for index, method in enumerate(sorted(by_method)):
        ramp = METHOD_RAMPS[index % len(METHOD_RAMPS)]  # wählt Farb Familie aus (für index=0 blau)
        entries = sorted(by_method[method], key=lambda item: (item[0] is None, item[0]))  #Sortiert (eps, Label)-Paare nach ε aufsteigend
        count = len(entries)
        for position, (_, label) in enumerate(entries):
            if count == 1:
                # Nur eine Variante: mittlere Farbe der Familie nehmen.
                shade = ramp[len(ramp) // 2]
            else:
                # Position des Labels auf Position der entsprechenden Farbe abstimmen
                step = round(position * (len(ramp) - 1) / (count - 1))
                shade = ramp[step]
            colors[label] = shade
    return colors

# mithilfe von Claude generiert
def method_chains(labels: List[str]) -> List[List[str]]:
    """Pfeilketten (vorher keine eigene Funktion): Baseline als eine, jedes Verfahren nach eps sortiert."""
    chains: List[List[str]] = [] # Jede innere Liste ist Kette von Labels

    # Baseline-Kette: real -> anon -> anon_synth
    baseline = [label for label in BASE_LABELS if label in labels]
    if len(baseline) > 1:
        chains.append(baseline)

    by_method: Dict[str, List[Tuple[Optional[float], str]]] = {}
    for label in labels:
        if label in BASE_COLORS: #Baseline überspringen
            continue
        method, eps = split_label(label)
        by_method.setdefault(method, []).append((eps, label))

    #Pro Verfahren eine Kette entlang von epsilon
    for method in sorted(by_method):
        entries = sorted(by_method[method], key=lambda item: (item[0] is None, item[0]))
        chain = [label for _, label in entries]  #behält nur label
        if len(chain) > 1:
            chains.append(chain)
    return chains


def as_float(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value))
    except ValueError:
        return None


def summary_value(row: Dict[str, Any], key: str) -> Optional[float]:
    return as_float(row.get(key))


def get_metric(row: Dict[str, Any], base_key: str, label: str, fallback_key: Optional[str] = None) -> Optional[float]:
    value = summary_value(row, f"{base_key}_{label}")
    if value is None and fallback_key:
        value = summary_value(row, f"{fallback_key}_{label}")
    return value


def compute_stats(values: List[float]) -> Tuple[Optional[float], Optional[float]]:
    if not values:
        return None, None
    if len(values) == 1:
        return values[0], 0.0
    return statistics.fmean(values), statistics.stdev(values)


def collect_metric(
    rows: List[Dict[str, Any]],
    metric: str,
    label: str,
    fallback_key: Optional[str] = None,
) -> List[float]:
    values: List[float] = []
    for row in rows:
        value = get_metric(row, metric, label, fallback_key=fallback_key)
        if value is not None:
            values.append(value)
    return values


def plot_privacy_utility_tradeoff(
    rows: List[Dict[str, Any]],
    output_path: Path,
    labels: List[str],
    colors: Dict[str, str],
) -> None:
    """x = Privacy, y = Utility, ein Punkt pro Variante.
    labels und colors werden uebergeben statt fest im Code zu stehen."""
    if plt is None:
        raise RuntimeError("matplotlib is required to plot figures. Install it with: pip install matplotlib")

    fig, ax = plt.subplots(figsize=(9, 6))
    # Mittelwerte pro Variante merken, damit die Pfeile spaeter daran ansetzen koennen.
    mean_points: Dict[str, Tuple[float, float]] = {}

    for label in labels:
        # Einzelne Runs blass im Hintergrund.
        points = []
        for row in rows:
            risk = get_metric(row, "privacy_score", label)
            utility = get_metric(row, "combined_utility", label, fallback_key="overall_utility")
            if risk is None or utility is None:
                continue
            points.append((risk, utility))
        if points:
            xs, ys = zip(*points)
            ax.scatter(
                xs,
                ys,
                s=14,
                color=colors[label],
                alpha=0.20,
                edgecolor="none",
                zorder=1,
            )

        # Mittelwert als grosser Punkt.
        risks = collect_metric(rows, "privacy_score", label)
        utilities = collect_metric(rows, "combined_utility", label, fallback_key="overall_utility")
        mean_risk, _ = compute_stats(risks)
        mean_utility, _ = compute_stats(utilities)
        if mean_risk is None or mean_utility is None:
            continue
        mean_points[label] = (mean_risk, mean_utility)
        ax.scatter(
            mean_risk,
            mean_utility,
            s=70,
            color=colors[label],
            edgecolor="white",
            linewidth=0.8,
            label=label,
            zorder=5,
        )

    # Pfeile nur innerhalb einer sinnvollen Ordnung:
    # Baseline-Pipeline und je Verfahren entlang steigender eps.
    for chain in method_chains(labels):
        # Paare aus aufeinanderfolgenden Labels: [a, b, c] -> (a, b), (b, c)
        for previous, current in zip(chain, chain[1:]):
            if previous in mean_points and current in mean_points:
                ax.annotate(
                    "",
                    xy=mean_points[current],
                    xytext=mean_points[previous],
                    arrowprops=dict(arrowstyle="->", color="#777777", alpha=0.55, lw=1.0),
                    zorder=3,
                )

    ax.set_title("Privacy-Utility Tradeoff")
    ax.set_xlabel("Privacy Score (1 - risk)")
    ax.set_ylabel("Combined Utility")
    ax.set_xlim(0.0, 1.0)
    ax.set_ylim(0.0, 1.0)
    ax.grid(True, linestyle="--", alpha=0.3)
    ax.legend(frameon=False, fontsize=8, loc="lower left")
    fig.tight_layout()
    fig.savefig(output_path, dpi=200)
    plt.close(fig)


def plot_summary_bars(
    rows: List[Dict[str, Any]],
    output_path: Path,
    labels: List[str],
    colors: Dict[str, str],
) -> None:
    """Zwei Balkendiagramme nebeneinander (Utility | Privacy), Mittelwert mit Std als Fehlerbalken."""
    if plt is None:
        raise RuntimeError("matplotlib is required to plot figures. Install it with: pip install matplotlib")

    # (Spalte, Ersatzspalte, Titel)
    metrics = [
        ("combined_utility", "overall_utility", "Combined Utility"),
        ("privacy_score", None, "Privacy Score (1 - risk)"),
    ]

    # Bild wird breiter, je mehr Varianten es gibt, damit die Balken nicht zu schmal werden.
    width = max(11.0, 1.1 * len(labels) * 2)
    fig, axes = plt.subplots(1, 2, figsize=(width, 5.0))
    for ax, (metric, fallback, title) in zip(axes, metrics):
        means = []
        stds = []
        for label in labels:
            values = collect_metric(rows, metric, label, fallback_key=fallback)
            mean, std = compute_stats(values)
            means.append(mean)
            stds.append(std)
        ax.bar(
            labels,
            means,
            yerr=stds,
            capsize=3,
            color=[colors[label] for label in labels],
            alpha=0.85,
        )
        ax.set_title(title)
        ax.set_ylim(0.0, 1.0)
        # Beschriftung schraeg stellen, sonst ueberlappen die langen Labels.
        ax.set_xticks(range(len(labels)))
        ax.set_xticklabels(labels, rotation=35, ha="right", fontsize=8)
        ax.grid(True, axis="y", linestyle="--", alpha=0.3)

    fig.tight_layout()
    fig.savefig(output_path, dpi=200)
    plt.close(fig)


def resolve_summary_path(root: Path, summary_path: Optional[str]) -> Path:
    if summary_path:
        return Path(summary_path)
    xlsx_path = root / SUMMARY_XLSX
    if xlsx_path.exists():
        return xlsx_path
    csv_path = root / SUMMARY_CSV
    if csv_path.exists():
        return csv_path
    raise FileNotFoundError(
        f"No summary file found. Expected {xlsx_path} or {csv_path}."
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Visualize evaluation summary results.")
    parser.add_argument(
        "--root",
        default=str(DEFAULT_ROOT),
        help="Base folder containing evaluation_summary.xlsx or .csv.",
    )
    parser.add_argument(
        "--summary",
        help="Explicit path to the summary file (xlsx or csv).",
    )
    # Optional: nur bestimmte Varianten plotten, z.B. --labels anon dpgan_eps1
    parser.add_argument(
        "--labels",
        nargs="*",
        help="Explicit labels to plot (default: auto-detect from columns).",
    )
    args = parser.parse_args()

    root = Path(args.root)
    summary_path = resolve_summary_path(root, args.summary)
    rows = read_summary(summary_path)
    if not rows:
        raise RuntimeError(f"No data found in {summary_path}")

    # Ohne --labels werden alle Varianten automatisch aus der Summary erkannt.
    labels = args.labels if args.labels else discover_labels(rows)
    colors = build_colors(labels)
    print(f"Labels: {', '.join(labels)}")

    output_dir = summary_path.parent
    plot_privacy_utility_tradeoff(
        rows, output_dir / "privacy_utility_tradeoff.png", labels, colors
    )
    plot_summary_bars(rows, output_dir / "summary_bars.png", labels, colors)
    print(f"Saved figures to {output_dir}")


if __name__ == "__main__":
    main()