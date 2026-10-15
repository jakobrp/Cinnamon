import argparse
import csv
import json
import math
import re
import statistics
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

try:
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    from matplotlib.colors import ListedColormap, Normalize
except ImportError:  # pragma: no cover - optional dependency
    plt = None
    Line2D = None
    ListedColormap = None
    Normalize = None

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - optional dependency
    load_workbook = None

try:
    import numpy as np
except ImportError:  # pragma: no cover - optional dependency
    np = None

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_ROOT = BASE_DIR / "experiment_results"

SUMMARY_XLSX = "evaluation_summary.xlsx"
SUMMARY_CSV = "evaluation_summary.csv"

# Feste Reihenfolge für die Baseline; alles Weitere wird erkannt.
BASE_LABELS = ["real", "anon", "anon_synth"]

# Anzeigenamen für die Baseline in Legende und Heatmap.
BASE_TITLES = {
    "real": "Real",
    "anon": "Anonymized",
    "anon_synth": "Anon+Synth",
}

# Feste Farben der Baseline, damit sie in allen Plots gleich aussehen.
BASE_COLORS = {
    "real": "#333333",
    "anon": "#1b6ca8",
    "anon_synth": "#2a9d8f",
}

# Farbfamilien pro Verfahren: hell (kleines eps) -> dunkel (grosses eps).
METHOD_RAMPS = [
    ["#9ecae1", "#4292c6", "#08519c"],  # Blau
    ["#fcae91", "#fb6a4a", "#a50f15"],  # Rot

]


def split_label(label: str) -> Tuple[str, Optional[float]]:
    """'dpgan_eps0_1' -> ('dpgan', 0.1). Ohne eps: (label, None)."""
    if "_eps" not in label:
        return label, None
    # Alles vor "_eps" ist das Verfahren, alles danach der eps-Wert.
    method, _, raw = label.partition("_eps")
    try:
        # Im Ordnernamen steht "_" statt ".", also 0_1 -> 0.1
        value = float(raw.replace("_", "."))
    except ValueError:
        value = None
    return method, value


def label_title(label: str) -> str:
    """Lesbarer Name für die Legende, z.B. 'dpgan_eps0_1' -> 'DPGAN eps=0.1'."""
    if label in BASE_TITLES:
        return BASE_TITLES[label]
    method, eps = split_label(label)
    pretty = method.upper().replace("_", " ")
    if eps is None:
        return pretty

    eps_text = f"{eps:g}"
    return f"{pretty} \u03b5={eps_text}"


def group_by_method(labels: List[str]) -> Dict[str, List[Tuple[Optional[float], str]]]:
    """Ausgelagert aus build_colors und method_chains"""
    by_method: Dict[str, List[Tuple[Optional[float], str]]] = {}
    for label in labels:
        if label in BASE_LABELS:
            continue
        method, eps = split_label(label)
        # Legt für ein neues Verfahren eine leere Liste an und hängt das Label an.
        by_method.setdefault(method, []).append((eps, label))
    for entries in by_method.values():
        # Nach eps sortieren; Labels ohne eps kommen ans Ende.
        entries.sort(key=lambda item: (item[0] is None, item[0]))
    return by_method

#Von claude generiert
def build_colors(labels: List[str]) -> Dict[str, str]:
    """Pro Verfahren eine Farbfamilie, innerhalb davon nach eps abgestuft."""
    colors = dict(BASE_COLORS)
    by_method = group_by_method(labels)

    # Jedes Verfahren bekommt eine eigene Farbfamilie (alphabetisch: dpgan -> Blau, pategan -> Rot).
    for index, method in enumerate(sorted(by_method)):
        ramp = METHOD_RAMPS[index % len(METHOD_RAMPS)]
        entries = by_method[method]
        count = len(entries)
        for position, (_, label) in enumerate(entries):
            if count == 1:
                # Nur eine Variante: mittlere Farbe nehmen.
                shade = ramp[len(ramp) // 2]
            else:
                # Varianten gleichmässig über die Farbfamilie verteilen.
                # Bei 3 eps-Werten: 0 -> hell, 1 -> mittel, 2 -> dunkel
                step = round(position * (len(ramp) - 1) / (count - 1))
                shade = ramp[step]
            colors[label] = shade
    return colors

#Von Claude generiert
def method_chains(labels: List[str]) -> List[List[str]]:
    """Ketten für die Pfeile: Baseline-Pipeline, dann je Verfahren nach eps."""
    chains: List[List[str]] = []

    # Baseline: real -> anon -> anon_synth
    baseline = [label for label in BASE_LABELS if label in labels]
    if len(baseline) > 1:
        chains.append(baseline)

    # Pro Verfahren: kleines eps -> grosses eps
    by_method = group_by_method(labels)
    for method in sorted(by_method):
        chain = [label for _, label in by_method[method]]
        if len(chain) > 1:
            chains.append(chain)
    return chains


def read_summary(summary_path: Path) -> List[Dict[str, Any]]:
    if summary_path.suffix.lower() == ".xlsx":
        return read_summary_xlsx(summary_path)
    return read_summary_csv(summary_path)


def read_summary_xlsx(summary_path: Path) -> List[Dict[str, Any]]:
    if load_workbook is None:
        raise RuntimeError("openpyxl is required to read Excel files. Install it with: pip install openpyxl")
    workbook = load_workbook(summary_path, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    data_rows: List[Dict[str, Any]] = []
    for row in rows[1:]:
        row_dict = {headers[i]: row[i] for i in range(len(headers))}
        data_rows.append(row_dict)
    return data_rows


def read_summary_csv(summary_path: Path) -> List[Dict[str, Any]]:
    data_rows: List[Dict[str, Any]] = []
    with summary_path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            data_rows.append(row)
    return data_rows


def discover_labels(dataset_rows: Dict[str, List[Dict[str, Any]]]) -> List[str]:
    """Vereinigung der Label über ALLE Datensätze, damit Farben konsistent sind."""
    # Jede Variante hat eine Spalte "privacy_score_<label>" in der Summary.
    prefix = "privacy_score_"
    found = set()
    for rows in dataset_rows.values():
        if not rows:
            continue
        for key in rows[0]:
            if key.startswith(prefix):
                found.add(key[len(prefix):])
    # Reihenfolge: erst Baseline, dann pro Verfahren nach eps sortiert.
    ordered = [label for label in BASE_LABELS if label in found]
    extras = sorted(found - set(BASE_LABELS))
    by_method = group_by_method(extras)
    for method in sorted(by_method):
        ordered.extend(label for _, label in by_method[method])
    return ordered


def as_float(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value))
    except ValueError:
        return None


def summary_value(row: Dict[str, Any], key: str) -> Optional[float]:
    return as_float(row.get(key))


def get_metric(row: Dict[str, Any], base_key: str, label: str, fallback_key: Optional[str] = None) -> Optional[float]:
    value = summary_value(row, f"{base_key}_{label}")
    if value is None and fallback_key:
        value = summary_value(row, f"{fallback_key}_{label}")
    return value


def compute_stats(values: List[float]) -> Tuple[Optional[float], Optional[float]]:
    if not values:
        return None, None
    if len(values) == 1:
        return values[0], 0.0
    return statistics.fmean(values), statistics.stdev(values)


def collect_metric(
    rows: List[Dict[str, Any]],
    metric: str,
    label: str,
    fallback_key: Optional[str] = None,
) -> List[float]:
    values: List[float] = []
    for row in rows:
        value = get_metric(row, metric, label, fallback_key=fallback_key)
        if value is not None:
            values.append(value)
    return values


def prettify_dataset_name(dataset_name: str) -> str:
    return re.sub(r"[_-]+", " ", dataset_name).strip().title()


def flatten_axes(axes):
    if isinstance(axes, list):
        return axes
    if hasattr(axes, "shape"):
        if len(axes.shape) == 0:
            return [axes]
        if len(axes.shape) == 1:
            return list(axes)
        return [ax for row in axes for ax in row]
    return [axes]


def format_dataset_title(dataset_name: str, size: Optional[int]) -> str:
    base = prettify_dataset_name(dataset_name)
    if size:
        return f"{base} (N={size:,})"
    return base


def resolve_dataset_csv(dataset_name: str) -> Optional[Path]:
    dataset_dir = BASE_DIR / "dataset_configs" / dataset_name
    if not dataset_dir.exists():
        return None
    preferred = dataset_dir / "original-dataset.csv"
    if preferred.exists():
        return preferred
    csv_files = sorted(dataset_dir.glob("*.csv"))
    if csv_files:
        return csv_files[0]
    return None


def get_dataset_size(dataset_name: str) -> Optional[int]:
    dataset_path = resolve_dataset_csv(dataset_name)
    if dataset_path is None:
        return None
    try:
        with dataset_path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
            reader = csv.reader(handle)
            rows = 0
            for _ in reader:
                rows += 1
        if rows <= 1:
            return 0
        return rows - 1
    except OSError:
        return None


def plot_privacy_utility_tradeoff(
    dataset_rows: Dict[str, List[Dict[str, Any]]],
    dataset_sizes: Dict[str, Optional[int]],
    output_path: Path,
    labels: List[str],
    colors: Dict[str, str],
) -> None:
    """Ein Teilplot pro Datensatz: x = Privacy, y = Utility, ein Punkt pro Variante."""
    if plt is None:
        raise RuntimeError("matplotlib is required to plot figures. Install it with: pip install matplotlib")

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.sans-serif": ["DejaVu Sans", "Arial", "Liberation Sans"],
            "axes.titlesize": 11,
            "axes.labelsize": 9,
            "xtick.labelsize": 8,
            "ytick.labelsize": 8,
            "axes.facecolor": "#fbfbfb",
            "figure.facecolor": "white",
            "grid.color": "#c9c9c9",
            "grid.alpha": 0.35,
            "grid.linewidth": 0.6,
        }
    )

    dataset_names = sorted(dataset_rows.keys())
    if not dataset_names:
        raise RuntimeError("No datasets found to plot.")

    ncols = 3 if len(dataset_names) > 3 else len(dataset_names)
    nrows = math.ceil(len(dataset_names) / ncols)
    fig, axes = plt.subplots(
        nrows=nrows,
        ncols=ncols,
        figsize=(4.6 * ncols, 3.6 * nrows),
        sharex=True,
        sharey=True,
    )
    axes_list = flatten_axes(axes)

    # Reihenfolgen für die Pfeile einmal vorab bestimmen (gilt für alle Datensätze)
    chains = method_chains(labels)

    for ax, dataset_name in zip(axes_list, dataset_names):
        rows = dataset_rows[dataset_name]
        # Mittelwerte pro Variante merken, damit die Pfeile später daran ansetzen können
        mean_points: Dict[str, Tuple[float, float]] = {}

        for label in labels:
            # Einzelne Runs blass im Hintergrund
            points: List[Tuple[float, float]] = []
            for row in rows:
                risk = get_metric(row, "privacy_score", label)
                utility = get_metric(row, "combined_utility", label, fallback_key="overall_utility")
                if risk is None or utility is None:
                    continue
                points.append((risk, utility))
            if points:
                xs, ys = zip(*points)
                ax.scatter(
                    xs,
                    ys,
                    s=14,
                    color=colors[label],
                    alpha=0.22,
                    edgecolor="none",
                    zorder=1,
                )

            # Mittelwert über alle Runs als grosser Punkt darüber
            risks = collect_metric(rows, "privacy_score", label)
            utilities = collect_metric(rows, "combined_utility", label, fallback_key="overall_utility")
            mean_risk, _ = compute_stats(risks)
            mean_utility, _ = compute_stats(utilities)
            if mean_risk is None or mean_utility is None:
                continue

            mean_points[label] = (mean_risk, mean_utility)
            ax.scatter(
                mean_risk,
                mean_utility,
                s=46,
                color=colors[label],
                edgecolor="white",
                linewidth=0.7,
                zorder=5,
            )

        # Pfeile: Baseline-Pipeline und je Verfahren entlang steigender eps.
        for chain in chains:
            # Paare aus aufeinanderfolgenden Labels: [a, b, c] -> (a, b), (b, c)
            for previous, current in zip(chain, chain[1:]):
                if previous in mean_points and current in mean_points:
                    ax.annotate(
                        "",
                        xy=mean_points[current],
                        xytext=mean_points[previous],
                        arrowprops=dict(arrowstyle="->", color="#555555", lw=1.0, alpha=0.7),
                        zorder=3,
                    )

        ax.set_title(format_dataset_title(dataset_name, dataset_sizes.get(dataset_name)))
        ax.set_xlim(0.0, 1.0)
        ax.set_ylim(0.0, 1.0)
        ax.grid(True, which="major", linestyle="-", alpha=0.35)
        ax.minorticks_on()
        ax.grid(True, which="minor", linestyle=":", alpha=0.18)
        ax.tick_params(direction="out", length=3, width=0.8, colors="#444444")
        ax.spines["top"].set_visible(True)
        ax.spines["right"].set_visible(True)
        ax.spines["top"].set_color("#111111")
        ax.spines["right"].set_color("#111111")
        ax.spines["left"].set_color("#666666")
        ax.spines["bottom"].set_color("#666666")

    for ax in axes_list[len(dataset_names):]:
        ax.axis("off")

    fig.supxlabel("Privacy Score", y=0.015)
    fig.supylabel("Utility Score")

    legend_handles = [
        Line2D(
            [0],
            [0],
            marker="o",
            color="none",
            markerfacecolor=colors[label],
            markeredgecolor="white",
            markeredgewidth=0.8,
            markersize=8,
            label=label_title(label),
        )
        for label in labels
    ]

    legend_cols = min(len(labels), 5)
    legend_rows = math.ceil(len(labels) / legend_cols)

    top = 0.93 - 0.02 * (legend_rows - 1)
    fig.legend(
        handles=legend_handles,
        loc="upper center",
        ncol=legend_cols,
        frameon=False,
        fontsize=8,
        bbox_to_anchor=(0.5, 0.99),
    )
    fig.tight_layout(rect=(0, 0.02, 1, top))
    fig.savefig(output_path, dpi=300, bbox_inches="tight", pad_inches=0.03)
    plt.close(fig)


def plot_privacy_utility_heatmap(
    dataset_rows: Dict[str, List[Dict[str, Any]]],
    output_path: Path,
    labels: List[str],
) -> None:
    """Zeilen = Varianten, Spaltengruppen = Datensätze (Utility | Privacy)."""
    # Gedreht gegenüber vorher, weil mit den DP-Varianten sonst zu viele Spalten entstehen.
    if plt is None or ListedColormap is None or Normalize is None:
        raise RuntimeError("matplotlib is required to plot figures. Install it with: pip install matplotlib")
    if np is None:
        raise RuntimeError("numpy is required to plot heatmaps. Install it with: pip install numpy")

    dataset_names = sorted(dataset_rows.keys())
    if not dataset_names:
        raise RuntimeError("No datasets found to plot.")

    # Pro Datensatz zwei Spalten: (Titel, Spaltenname in der Summary, Ersatzspalte)
    metrics = [
        ("Utility", "combined_utility", "overall_utility"),
        ("Privacy", "privacy_score", None),
    ]
    # Breite der Lücke zwischen zwei Datensätzen.
    gap_width = 0.5

    # Spalten festlegen: Datensatz A (Utility, Privacy), Lücke, Datensatz B (...), ...
    col_specs: List[Dict[str, Any]] = []
    for index, dataset_name in enumerate(dataset_names):
        for metric_title, metric_key, fallback in metrics:
            col_specs.append(
                {
                    "dataset": dataset_name,
                    "metric_title": metric_title,
                    "metric_key": metric_key,
                    "fallback": fallback,
                    "is_gap": False,
                }
            )
        if index < len(dataset_names) - 1:
            col_specs.append({"is_gap": True})

    # Matrix fuellen: eine Zeile pro Variante, eine Spalte pro col_spec.
    data: List[List[float]] = []
    stds: List[List[Optional[float]]] = []
    for label in labels:
        row_values: List[float] = []
        row_stds: List[Optional[float]] = []
        for spec in col_specs:
            if spec["is_gap"]:
                # NaN wird später weiss gezeichnet.
                row_values.append(float("nan"))
                row_stds.append(None)
                continue
            values = collect_metric(
                dataset_rows[spec["dataset"]],
                spec["metric_key"],
                label,
                fallback_key=spec["fallback"],
            )
            mean, std = compute_stats(values)
            row_values.append(mean if mean is not None else float("nan"))
            row_stds.append(std if mean is not None else None)
        data.append(row_values)
        stds.append(row_stds)

    # Zahlen zusätzlich als CSV und JSON ablegen (Grundlage für Tabellen).
    # Long-Format: eine Zeile pro Datensatz x Variante x Metrik,
    # damit neue Varianten keine neuen Spalten erzeugen.
    csv_output = output_path.with_suffix(".csv")
    json_output = output_path.with_suffix(".json")

    with csv_output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["dataset", "variant", "metric", "mean", "std"])
        for label, row_values, row_stds in zip(labels, data, stds):
            for spec, value, std in zip(col_specs, row_values, row_stds):
                # value != value ist nur bei NaN wahr -> Lücken und fehlende Werte überspringen.
                if spec["is_gap"] or value != value:
                    continue
                writer.writerow(
                    [spec["dataset"], label, spec["metric_title"].lower(), value, std]
                )

    json_rows: List[Dict[str, Any]] = []
    for label, row_values, row_stds in zip(labels, data, stds):
        for spec, value, std in zip(col_specs, row_values, row_stds):
            if spec["is_gap"] or value != value:
                continue
            json_rows.append(
                {
                    "dataset": spec["dataset"],
                    "variant": label,
                    "metric": spec["metric_title"].lower(),
                    "mean": value,
                    "std": std,
                }
            )
    with json_output.open("w", encoding="utf-8") as handle:
        json.dump(json_rows, handle, indent=2)

    base = plt.get_cmap("YlGnBu")
    ramp = base(np.linspace(0.82, 0, 256))
    cmap = ListedColormap(ramp)
    cmap.set_bad("#ffffff")
    norm = Normalize(vmin=0, vmax=1)

    # Bildgrösse wächst mit der Anzahl der Spalten (Breite) und Varianten (Höhe).
    col_widths = [gap_width if spec["is_gap"] else 1.0 for spec in col_specs]
    total_width = sum(col_widths)
    fig_width = max(8.0, 0.95 * total_width + 3.0)
    fig_height = max(3.0, 0.48 * len(labels) + 1.8)
    fig, ax = plt.subplots(figsize=(fig_width, fig_height))

    x_edges = [0.0]
    for width in col_widths:
        x_edges.append(x_edges[-1] + width)
    y_edges = list(range(len(labels) + 1))

    im = ax.pcolormesh(x_edges, y_edges, data, cmap=cmap, norm=norm, shading="flat")
    ax.invert_yaxis()

    x_tick_positions = []
    x_tick_labels = []
    centers = []
    for spec, width, start in zip(col_specs, col_widths, x_edges[:-1]):
        center = start + width / 2
        centers.append(center)
        if not spec["is_gap"]:
            x_tick_positions.append(center)
            x_tick_labels.append(spec["metric_title"])

    ax.set_xticks(x_tick_positions)
    ax.set_xticklabels(x_tick_labels, fontsize=8)
    ax.set_yticks([y + 0.5 for y in range(len(labels))])
    ax.set_yticklabels([label_title(label) for label in labels], fontsize=9)
    ax.tick_params(axis="x", pad=6)

    ax.set_xticks(x_edges, minor=True)
    ax.set_yticks(y_edges, minor=True)
    ax.grid(which="minor", color="#e1e1e1", linewidth=0.6)
    ax.tick_params(which="minor", bottom=False, left=False)

    # Datensatznamen über die Spaltengruppen setzen.
    col_index = 0
    for index, dataset_name in enumerate(dataset_names):
        first = col_index
        last = col_index + len(metrics) - 1
        center = (x_edges[first] + x_edges[last + 1]) / 2
        ax.text(
            center,
            1.03,
            prettify_dataset_name(dataset_name),
            transform=ax.get_xaxis_transform(),
            ha="center",
            va="bottom",
            fontsize=9,
            color="#222222",
        )
        col_index += len(metrics)
        if index < len(dataset_names) - 1:
            col_index += 1  # Luecke überspringen

    # Mittelwert +- Standardabweichung in jede Zelle schreiben.
    for i in range(len(labels)):
        for j, center in enumerate(centers):
            value = data[i][j]
            std = stds[i][j]
            if value != value:
                continue
            text = f"{value:.3f}" if std is None else f"{value:.3f}\u00b1{std:.3f}"
            ax.text(
                center,
                i + 0.5,
                text,
                ha="center",
                va="center",
                fontsize=7.5,
                color="#111111",
            )

    cbar = fig.colorbar(im, ax=ax, fraction=0.025, pad=0.02)
    cbar.ax.set_ylabel("Score", rotation=90, labelpad=8)
    cbar.ax.tick_params(labelsize=8)

    ax.set_xlabel("")
    ax.set_ylabel("")
    ax.spines["top"].set_visible(False)
    fig.tight_layout(rect=(0, 0, 1, 0.94))
    fig.savefig(output_path, dpi=300, bbox_inches="tight", pad_inches=0.03)
    plt.close(fig)


def resolve_summary_path(root: Path, summary_path: Optional[str]) -> Path:
    if summary_path:
        return Path(summary_path)
    xlsx_path = root / SUMMARY_XLSX
    if xlsx_path.exists():
        return xlsx_path
    csv_path = root / SUMMARY_CSV
    if csv_path.exists():
        return csv_path
    raise FileNotFoundError(f"No summary file found. Expected {xlsx_path} or {csv_path}.")


def collect_dataset_summaries(root: Path, summary_path: Optional[str]) -> Dict[str, Path]:
    if summary_path:
        summary = Path(summary_path)
        return {summary.parent.name: summary}

    if (root / SUMMARY_XLSX).exists() or (root / SUMMARY_CSV).exists():
        return {root.name: resolve_summary_path(root, None)}

    datasets: Dict[str, Path] = {}
    for child in sorted(root.iterdir()):
        if not child.is_dir():
            continue
        try:
            datasets[child.name] = resolve_summary_path(child, None)
        except FileNotFoundError:
            continue
    return datasets


def main() -> None:
    parser = argparse.ArgumentParser(description="Visualize evaluation summary results.")
    parser.add_argument(
        "--root",
        default=str(DEFAULT_ROOT),
        help="Base folder containing dataset subfolders with evaluation_summary.xlsx or .csv.",
    )
    parser.add_argument(
        "--summary",
        help="Explicit path to the summary file (xlsx or csv).",
    )
    # Optional: nur bestimmte Varianten plotten, z.B. --labels anon dpgan_eps1
    parser.add_argument(
        "--labels",
        nargs="*",
        help="Explicit variants to plot (default: auto-detect from columns).",
    )
    args = parser.parse_args()

    root = Path(args.root)
    if not root.exists():
        raise FileNotFoundError(f"Root folder not found: {root}")

    dataset_summaries = collect_dataset_summaries(root, args.summary)
    if not dataset_summaries:
        raise RuntimeError(f"No evaluation summaries found under {root}")

    dataset_rows: Dict[str, List[Dict[str, Any]]] = {}
    dataset_sizes: Dict[str, Optional[int]] = {}
    for dataset_name, summary_path in dataset_summaries.items():
        rows = read_summary(summary_path)
        if not rows:
            continue
        dataset_rows[dataset_name] = rows
        dataset_sizes[dataset_name] = get_dataset_size(dataset_name)

    if not dataset_rows:
        raise RuntimeError("No data found in any evaluation summary.")

    # Ohne --labels werden alle Varianten automatisch aus den Summaries erkannt.
    labels = args.labels if args.labels else discover_labels(dataset_rows)
    colors = build_colors(labels)
    print(f"Datasets: {', '.join(sorted(dataset_rows))}")
    print(f"Variants: {', '.join(labels)}")

    output_dir = root
    plot_privacy_utility_tradeoff(
        dataset_rows,
        dataset_sizes,
        output_dir / "privacy_utility_tradeoff_overview.png",
        labels,
        colors,
    )
    plot_privacy_utility_heatmap(
        dataset_rows,
        output_dir / "privacy_utility_heatmap_overview.png",
        labels,
    )
    print(f"Saved figures to {output_dir}")


if __name__ == "__main__":
    main()